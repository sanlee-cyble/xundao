#!/usr/bin/env python3
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


root = Path(sys.argv[1]).expanduser().resolve()
root.mkdir(parents=True, exist_ok=True)

dynamic = Workbook()
dynamic_sheet = dynamic.active
dynamic_sheet.title = "动态任务"
dynamic_sheet.append([
    "名称",
    "蒲公英链接",
    "粉丝数（w）",
    "报价",
    "预估合作笔记自然流阅读（90天）",
    "下单价",
    "实际花费",
    "CPC",
])
dynamic_sheet.append([
    "测试达人",
    "https://pgy.xiaohongshu.com/solar/pre-trade/blogger-detail/creator-fixture",
    "",
    "",
    "",
    "",
    "",
    "",
])
for cell in dynamic_sheet[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="DDEBF7")
dynamic_sheet.freeze_panes = "A2"
dynamic.save(root / "dynamic.xlsx")

template_only = Workbook()
template_only_sheet = template_only.active
template_only_sheet.title = "字段模板"
template_only_sheet.append([
    "达人信息",
    "",
    "90天合作笔记",
    "",
])
template_only_sheet.append([
    "名称",
    "粉丝数（w）",
    "曝光中位数（90天）",
    "阅读中位数（90天）",
])
for row in template_only_sheet.iter_rows(min_row=1, max_row=2):
    for cell in row:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E2F0D9")
template_only_sheet.freeze_panes = "A3"
template_only.save(root / "template-only.xlsx")

links_only = Workbook()
links_only_sheet = links_only.active
links_only_sheet.title = "达人链接"
links_only_sheet.append(["蒲公英链接"])
links_only_sheet.append([
    "https://pgy.xiaohongshu.com/solar/pre-trade/blogger-detail/creator-links-only",
])
links_only_sheet["A1"].font = Font(bold=True)
links_only_sheet["A1"].fill = PatternFill("solid", fgColor="FFF2CC")
links_only.save(root / "links-only.xlsx")

medela = Workbook()
medela_sheet = medela.active
medela_sheet.title = "KOL提报表"
medela_sheet.append([
    "提报日期",
    "序号",
    "达人类型",
    "名称",
    "推荐理由",
    "内容方向",
    "是否合作",
    "合作方式",
    "蒲公英链接",
    "主页链接",
    "粉丝数（w）",
    "合作笔记曝光中位数（90天）",
])
medela_sheet.append([""] * 12)
medela_sheet.append([
    "",
    1,
    "母婴",
    "测试达人",
    "",
    "",
    "",
    "",
    "https://pgy.xiaohongshu.com/solar/pre-trade/blogger-detail/creator-fixture",
])
for cell in medela_sheet[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="FCE4D6")
medela.save(root / "medela.xlsx")
