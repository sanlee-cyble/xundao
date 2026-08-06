#!/usr/bin/env python3
import argparse
import json
import re
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]


TARGET_FIELDS = [
    ("粉丝数", ["粉丝数"]),
    ("总互动数", ["总互动数（获赞与收藏）", "总互动数"]),
    ("图文报价", ["图文笔记（元）", "图文笔记"]),
    ("视频报价", ["视频笔记（元）", "视频笔记"]),
    ("预估阅读单价", ["预估阅读单价（元/阅读）", "预估阅读单价"]),
    ("预估互动单价", ["预估互动单价（元/互动）", "预估互动单价"]),
    ("阅读中位数", ["阅读中位数", "近30日阅读中位数"]),
    ("互动中位数", ["互动中位数", "近30日互动中位数"]),
    ("视频完播率", ["视频完播率"]),
    ("图文3秒阅读率", ["图文3秒阅读率"]),
    ("活跃粉丝占比", ["活跃粉丝占比"]),
    ("阅读粉丝占比", ["阅读粉丝占比"]),
    ("互动粉丝占比", ["互动粉丝占比"]),
    ("下单粉丝占比", ["下单粉丝占比"]),
    ("25岁以上粉丝占比", ["25岁以上粉丝占比"]),
    ("粉丝地域分布前3位", ["粉丝地域分部前3位", "粉丝地域分布前3位"]),
]

EXPORT_HEADERS = [
    "账号昵称",
    "小红书号",
    "蒲公英链接",
    "蒲公英达人ID",
    "数据状态",
    "失败原因",
    "数据更新至",
    *[field for field, _labels in TARGET_FIELDS],
    "证据文件",
]


def selected_target_fields(raw_fields=""):
    if not raw_fields:
        return TARGET_FIELDS
    try:
        selected = json.loads(raw_fields)
    except json.JSONDecodeError:
        selected = [item.strip() for item in str(raw_fields).split(",")]
    selected_set = {str(item).strip() for item in selected if str(item).strip()}
    if not selected_set:
        return TARGET_FIELDS
    picked = [item for item in TARGET_FIELDS if item[0] in selected_set]
    return picked or TARGET_FIELDS


def clean(value):
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def cell_text(sheet, row, col):
    value = sheet.cell(row=row, column=col).value
    if value is not None:
        return clean(value)
    for merged in sheet.merged_cells.ranges:
        if sheet.cell(row=row, column=col).coordinate in merged:
            return clean(sheet.cell(row=merged.min_row, column=merged.min_col).value)
    return ""


def col_header_text(sheet, col, header_end):
    labels = [cell_text(sheet, row, col) for row in range(1, header_end + 1)]
    return " ".join(label for label in labels if label)


def find_header_end(sheet):
    header_end = 1
    markers = {"账号昵称", "小红书号", "蒲公英链接", "图文笔记", "视频笔记", "预估阅读单价"}
    for row in range(1, min(sheet.max_row, 8) + 1):
        row_text = " ".join(cell_text(sheet, row, col) for col in range(1, sheet.max_column + 1))
        if any(marker in row_text for marker in markers):
            header_end = row
    return header_end


def find_col(sheet, header_end, *keywords):
    for col in range(1, sheet.max_column + 1):
        header = col_header_text(sheet, col, header_end)
        if all(keyword in header for keyword in keywords):
            return col
    return None


def find_col_by_labels(sheet, header_end, labels):
    for label in labels:
        col = find_col(sheet, header_end, label)
        if col:
            return col
    return None


def creator_id_from_url(url):
    match = re.search(r"/blogger-detail/([^/?#]+)", url or "")
    return match.group(1) if match else ""


def extract(args):
    input_path = Path(args.input).expanduser().resolve()
    wb = load_workbook(input_path, data_only=False)
    sheet = wb[args.sheet] if args.sheet else wb.active
    header_end = find_header_end(sheet)

    nickname_col = find_col(sheet, header_end, "账号昵称")
    red_id_col = find_col(sheet, header_end, "小红书号")
    link_col = find_col(sheet, header_end, "蒲公英链接")
    if link_col is None:
        raise RuntimeError("无法在表头中定位“蒲公英链接”列")

    creators = []
    for row in range(header_end + 1, sheet.max_row + 1):
        nickname = clean(sheet.cell(row=row, column=nickname_col).value) if nickname_col else ""
        red_id = clean(sheet.cell(row=row, column=red_id_col).value) if red_id_col else ""
        pgy_link = clean(sheet.cell(row=row, column=link_col).value)
        if not nickname and not red_id and not pgy_link:
            continue
        if "pgy.xiaohongshu.com" not in pgy_link:
            continue
        creators.append(
            {
                "rowIndex": row,
                "nickname": nickname,
                "redId": red_id,
                "pgyLink": pgy_link,
                "creatorId": creator_id_from_url(pgy_link),
            }
        )

    output = Path(args.creators).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(creators, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"input": str(input_path), "sheet": sheet.title, "headerEnd": header_end, "count": len(creators), "creators": str(output)}, ensure_ascii=False))


def is_blank(value):
    return value is None or value == ""


def normalize_number(value):
    if value is None or value == "":
        return ""
    if isinstance(value, (int, float)):
        return value
    text = str(value).replace(",", "").replace("¥", "").strip()
    if text == "-" or text == "":
        return text
    try:
        if "." in text:
            return float(text)
        return int(text)
    except ValueError:
        return value


def build_target_cols(sheet, header_end, target_fields=None):
    target_fields = target_fields or TARGET_FIELDS
    targets = {}
    for field, labels in target_fields:
        col = find_col_by_labels(sheet, header_end, labels)
        if col:
            targets[field] = col
    return targets


def field_value(fields, field):
    aliases = {
        "阅读中位数": ["阅读中位数", "近30日阅读中位数"],
        "互动中位数": ["互动中位数", "近30日互动中位数"],
    }.get(field, [field])
    for alias in aliases:
        value = fields.get(alias)
        if isinstance(value, dict):
            value = value.get("value")
        if value not in (None, ""):
            return value
    return ""


def copy_column_styles(sheet, header_end, target_cols, row):
    for col in target_cols.values():
        source = sheet.cell(row=header_end + 1, column=col)
        target = sheet.cell(row=row, column=col)
        if source.has_style and not target.has_style:
            target._style = copy(source._style)


def fill(args):
    input_path = Path(args.input).expanduser().resolve()
    output_path = Path(args.output).expanduser().resolve()
    results_path = Path(args.results).expanduser().resolve()
    results = json.loads(results_path.read_text(encoding="utf-8"))
    target_fields = selected_target_fields(args.fields)
    by_row = {int(row["rowIndex"]): row for row in results if row.get("rowIndex")}

    wb = load_workbook(input_path)
    sheet = wb[args.sheet] if args.sheet else wb.active
    header_end = find_header_end(sheet)
    link_col = find_col(sheet, header_end, "蒲公英链接")
    target_cols = build_target_cols(sheet, header_end, target_fields)

    filled = 0
    skipped_existing = 0
    missing_cols = [field for field, _labels in target_fields if field not in target_cols]
    for row_index, result in by_row.items():
        fields = result.get("fields") or {}
        copy_column_styles(sheet, header_end, target_cols, row_index)
        for field, col in target_cols.items():
            if link_col and col <= link_col:
                continue
            cell = sheet.cell(row=row_index, column=col)
            if not is_blank(cell.value):
                skipped_existing += 1
                continue
            value = field_value(fields, field)
            value = normalize_number(value)
            if value != "":
                cell.value = value
                filled += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(json.dumps({"output": str(output_path), "filled": filled, "skippedExisting": skipped_existing, "missingColumns": missing_cols}, ensure_ascii=False))


def export_results(args):
    output_path = Path(args.output).expanduser().resolve()
    results_path = Path(args.results).expanduser().resolve()
    results = json.loads(results_path.read_text(encoding="utf-8"))
    target_fields = selected_target_fields(args.fields)

    wb = Workbook()
    sheet = wb.active
    sheet.title = "达人表现数据"
    export_headers = [
        "账号昵称",
        "小红书号",
        "蒲公英链接",
        "蒲公英达人ID",
        "数据状态",
        "失败原因",
        "数据更新至",
        *[field for field, _labels in target_fields],
        "证据文件",
    ]
    sheet.append(export_headers)

    for result in results:
        fields = result.get("fields") or {}
        row = [
            clean(result.get("nickname")),
            clean(result.get("redId")),
            clean(result.get("pgyLink")),
            clean(result.get("creatorId")),
            clean(result.get("dataStatus")),
            clean(result.get("errorReason")),
            clean(result.get("dataUpdateDate")),
        ]
        for field, _labels in target_fields:
            row.append(normalize_number(field_value(fields, field)))
        row.append("；".join(result.get("evidenceFiles") or []))
        sheet.append(row)

    for cell in sheet[1]:
        cell.style = "Headline 4"
    sheet.freeze_panes = "A2"
    widths = {
        "A": 18,
        "B": 16,
        "C": 58,
        "D": 28,
        "E": 12,
        "F": 24,
        "G": 16,
        "W": 56,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(json.dumps({"output": str(output_path), "rows": len(results)}, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(description="蒲公英采集器 Excel 读写工具")
    sub = parser.add_subparsers(dest="command", required=True)

    extract_parser = sub.add_parser("extract")
    extract_parser.add_argument("--input", required=True)
    extract_parser.add_argument("--sheet", default="")
    extract_parser.add_argument("--creators", default=str(ROOT / "data/prototype_creators.json"))
    extract_parser.set_defaults(func=extract)

    fill_parser = sub.add_parser("fill")
    fill_parser.add_argument("--input", required=True)
    fill_parser.add_argument("--output", required=True)
    fill_parser.add_argument("--results", default=str(ROOT / "data/prototype_results.json"))
    fill_parser.add_argument("--sheet", default="")
    fill_parser.add_argument("--fields", default="")
    fill_parser.set_defaults(func=fill)

    export_parser = sub.add_parser("export")
    export_parser.add_argument("--output", required=True)
    export_parser.add_argument("--results", default=str(ROOT / "data/prototype_results.json"))
    export_parser.add_argument("--fields", default="")
    export_parser.set_defaults(func=export_results)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
