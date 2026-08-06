#!/usr/bin/env python3
import argparse
import html
import json
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


FORMULA_ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")


def json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def load_book(input_path):
    path = Path(input_path).expanduser().resolve()
    return load_workbook(
        path,
        data_only=False,
        keep_vba=path.suffix.lower() == ".xlsm",
        keep_links=True,
    )


def cell_value(cell):
    target = getattr(cell.hyperlink, "target", "") if cell.hyperlink else ""
    if "pgy.xiaohongshu.com" in str(target):
        return str(target)
    return json_value(cell.value)


def worksheet_values(sheet):
    rows = []
    last_nonempty = 0
    for row_index, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column),
        start=1,
    ):
        values = [cell_value(cell) for cell in row]
        while values and values[-1] in (None, ""):
            values.pop()
        if any(value not in (None, "") for value in values):
            last_nonempty = row_index
        rows.append(values)
    return rows[:last_nonempty]


def inspect_workbook(args):
    workbook = load_book(args.input)
    selected = workbook.worksheets[0] if workbook.worksheets else None
    selected_values = []
    for sheet in workbook.worksheets:
        values = worksheet_values(sheet)
        if not selected_values:
            selected = sheet
            selected_values = values
        if any(
            "pgy.xiaohongshu.com" in str(value or "")
            for row in values
            for value in row
        ):
            selected = sheet
            selected_values = values
            break
    if selected is None:
        raise RuntimeError("Excel 中没有工作表")
    print(json.dumps({
        "sheetName": selected.title,
        "values": selected_values,
    }, ensure_ascii=False))


def copy_cell_style(source, target):
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def best_template_sheet(workbook):
    selected = workbook.worksheets[0] if workbook.worksheets else None
    best_score = -1
    terms = ("蒲公英", "达人", "账号", "名称", "粉丝", "曝光", "阅读", "点赞", "收藏", "报价")
    for sheet in workbook.worksheets:
        score = 0
        for row in sheet.iter_rows(min_row=1, max_row=min(max(sheet.max_row, 1), 10)):
            for cell in row:
                value = str(cell.value or "")
                if any(term in value for term in terms):
                    score += 1
        if score > best_score:
            selected = sheet
            best_score = score
    return selected


def attach_links(args):
    links = json.loads(args.links)
    if not isinstance(links, list):
        raise RuntimeError("外部蒲公英链接格式无效")
    links = [
        str(value).strip()
        for value in links
        if "pgy.xiaohongshu.com" in str(value or "")
    ]
    if not links:
        raise RuntimeError("未找到可写入 Excel 的蒲公英链接")
    workbook = load_book(args.input)
    sheet = best_template_sheet(workbook)
    if sheet is None:
        raise RuntimeError("Excel 中没有工作表")

    header_terms = ("蒲公英", "达人", "账号", "名称", "粉丝", "曝光", "阅读", "点赞", "收藏", "报价", "CPC")
    header_row, link_column = find_header_column(sheet, "蒲公英链接", max_rows=10)
    scan_limit = min(max(sheet.max_row, 1), 10)
    scored_rows = []
    for row_index in range(1, scan_limit + 1):
        values = [str(sheet.cell(row=row_index, column=column).value or "") for column in range(1, sheet.max_column + 1)]
        score = sum(1 for value in values if any(term in value for term in header_terms))
        if score:
            scored_rows.append((score, row_index))
    inferred_header_row = max(scored_rows, key=lambda item: (item[0], item[1]))[1] if scored_rows else 1
    header_row = header_row or inferred_header_row
    header_end = max(
        [row_index for _, row_index in scored_rows if row_index >= header_row] or [header_row]
    )

    if not link_column:
        nonempty_columns = [
            column
            for row_index in range(1, header_end + 1)
            for column in range(1, sheet.max_column + 1)
            if not is_blank(sheet.cell(row=row_index, column=column).value)
        ]
        link_column = (max(nonempty_columns) if nonempty_columns else sheet.max_column) + 1
        header_cell = sheet.cell(row=header_end, column=link_column)
        header_cell.value = "蒲公英链接"
        style_source_column = max(link_column - 1, 1)
        copy_cell_style(sheet.cell(row=header_end, column=style_source_column), header_cell)

    first_data_row = header_end + 1
    for offset, link in enumerate(links):
        row_index = first_data_row + offset
        cell = sheet.cell(row=row_index, column=link_column)
        if row_index > first_data_row:
            copy_cell_style(sheet.cell(row=first_data_row, column=link_column), cell)
        cell.value = excel_safe_text(link)

    output = Path(args.output or args.input).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(json.dumps({
        "ok": True,
        "outputPath": str(output),
        "sheetName": sheet.title,
        "headerRow": header_end,
        "linkColumn": link_column,
        "count": len(links),
    }, ensure_ascii=False))


def is_blank(value):
    return value is None or str(value).strip() == ""


def excel_safe_text(value):
    text = str(value)
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def numeric_or_blank(value):
    if value is None or value == "":
        return ""
    if isinstance(value, (int, float)):
        return value
    text = str(value).replace(",", "").replace("¥", "").strip()
    if text.endswith("%"):
        try:
            return float(text[:-1]) / 100
        except ValueError:
            return ""
    try:
        number = float(text)
        return int(number) if number.is_integer() else number
    except ValueError:
        return excel_safe_text(value)


def set_if_blank(sheet, row, column, value):
    if is_blank(value):
        return False
    cell = sheet.cell(row=row, column=column)
    if not is_blank(cell.value):
        return False
    cell.value = value
    return True


def formula_for(field_id, row, by_field_id, task=None):
    quote = by_field_id.get("pgy.creator.video_quote")
    fans = by_field_id.get("pgy.creator.fans")
    order_price = by_field_id.get("derived.order_price")
    natural_read = by_field_id.get("pgy.coop.video.90d.natural.read")
    calculation_context = (task or {}).get("calculationContext") or {}
    exchange_rate_cell = str(calculation_context.get("exchangeRateCell") or "").strip()
    exchange_rate_match = re.fullmatch(r"([A-Z]+)(\d+)", exchange_rate_cell, re.IGNORECASE)
    if field_id == "derived.cost_jpy" and quote and exchange_rate_match:
        absolute_rate = f"${exchange_rate_match.group(1).upper()}${exchange_rate_match.group(2)}"
        return f'=IF({quote}{row}="","",ROUND({quote}{row}*{absolute_rate},2))'
    if field_id == "derived.cost_per_fan" and quote and fans:
        return f'=IF({fans}{row}>0,{quote}{row}/{fans}{row},"")'
    if field_id == "derived.order_price" and quote:
        return f'=IF({quote}{row}="","",ROUND({quote}{row}*1.1,2))'
    if field_id == "derived.actual_spend" and order_price:
        return f'=IF({order_price}{row}="","",ROUND({order_price}{row}*1.02,2))'
    if field_id == "derived.cpc" and quote and natural_read:
        return f'=IF({natural_read}{row}>0,{quote}{row}/{natural_read}{row},"")'
    return ""


def apply_dynamic_format(cell, mapping):
    field_id = (mapping or {}).get("id", "")
    unit = (mapping or {}).get("unit", "")
    if unit == "%":
        cell.number_format = "0.0%"
    elif field_id == "pgy.creator.fans_w":
        cell.number_format = "0.0000"
    elif field_id in (
        "pgy.creator.video_quote",
        "derived.cost_jpy",
        "derived.order_price",
        "derived.actual_spend",
    ):
        cell.number_format = "¥#,##0.00"
    elif field_id in ("derived.cpc", "derived.cost_per_fan"):
        cell.number_format = "0.00"
    elif unit in ("次", "人"):
        cell.number_format = "#,##0"


def value_for_dynamic_column(column, result):
    mapping = column.get("mapping") or {}
    if mapping.get("id") == "pgy.creator.pgy_profile":
        return result.get("pgyLink", "")
    contract_key = column.get("contractKey", "")
    contract_values = result.get("contractValues") or {}
    if contract_key and contract_key in contract_values:
        return contract_values[contract_key]
    return (result.get("fields") or {}).get(mapping.get("name", ""), "")


def formulas_and_errors(sheet):
    formulas = []
    errors = []
    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                formulas.append({"cell": cell.coordinate, "formula": value})
            if isinstance(value, str) and any(token in value for token in FORMULA_ERRORS):
                errors.append({"cell": cell.coordinate, "value": value})
    return formulas, errors


def table_snapshot(sheet, max_rows=50, max_cols=40):
    rows = []
    for row in sheet.iter_rows(
        min_row=1,
        max_row=min(sheet.max_row, max_rows),
        min_col=1,
        max_col=min(sheet.max_column, max_cols),
    ):
        rows.append([json_value(cell.value) for cell in row])
    return rows


def render_preview_svg(sheet, output_path, max_rows=22, max_cols=20):
    row_count = min(max(sheet.max_row, 1), max_rows)
    column_count = min(max(sheet.max_column, 1), max_cols)
    cell_width = 150
    row_height = 34
    title_height = 48
    width = column_count * cell_width
    height = title_height + row_count * row_height
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<rect width="100%" height="100%" fill="#f8fafc"/>',
        f'<text x="16" y="30" font-family="Arial, sans-serif" font-size="18" font-weight="700" fill="#152238">{html.escape(sheet.title)}</text>',
    ]
    for row_index in range(1, row_count + 1):
        y = title_height + (row_index - 1) * row_height
        fill = "#e8eef7" if row_index <= 2 else ("#ffffff" if row_index % 2 else "#f6f8fb")
        for column_index in range(1, column_count + 1):
            x = (column_index - 1) * cell_width
            value = sheet.cell(row=row_index, column=column_index).value
            text = str(value if value is not None else "").replace("\n", " ")
            if len(text) > 22:
                text = f"{text[:21]}…"
            parts.append(
                f'<rect x="{x}" y="{y}" width="{cell_width}" height="{row_height}" fill="{fill}" stroke="#cbd5e1"/>'
            )
            parts.append(
                f'<text x="{x + 8}" y="{y + 22}" font-family="Arial, sans-serif" font-size="12" fill="#24364b">{html.escape(text)}</text>'
            )
    parts.append("</svg>")
    output = Path(output_path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("".join(parts), encoding="utf-8")


def write_validation(path, payload):
    output = Path(path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def dynamic_fill(args):
    task = json.loads(Path(args.task).read_text(encoding="utf-8"))
    results = json.loads(Path(args.results).read_text(encoding="utf-8"))
    workbook = load_book(args.input)
    sheet = workbook[task["sheetName"]] if task.get("sheetName") in workbook.sheetnames else workbook.active
    header_depth = max(int(task.get("headerDepth") or 1), 1)
    appended_columns = []
    for column in task.get("columns", []):
        column_index = int(column["index"]) + 1
        if not column.get("appendedByTemplate"):
            continue
        if header_depth > 1:
            header_values = [
                (1, column.get("topLabel") or column.get("parentLabel") or ""),
                (header_depth, column.get("childLabel") or column.get("displayLabel") or ""),
            ]
        else:
            header_values = [(1, column.get("displayLabel") or column.get("childLabel") or "")]
        for row_index, value in header_values:
            cell = sheet.cell(row=row_index, column=column_index)
            if is_blank(cell.value):
                style_source = sheet.cell(row=row_index, column=max(column_index - 1, 1))
                copy_cell_style(style_source, cell)
                cell.value = excel_safe_text(value)
        appended_columns.append({
            "column": get_column_letter(column_index),
            "fieldId": (column.get("mapping") or {}).get("id", ""),
            "name": column.get("displayLabel") or "",
        })
    by_row = {int(item["rowIndex"]): item for item in results if item.get("rowIndex")}
    by_link = {str(item.get("pgyLink") or ""): item for item in results}
    by_field_id = {
        column["mapping"]["id"]: column["letter"]
        for column in task.get("columns", [])
        if column.get("mapping")
    }
    filled = []
    left_blank = []

    for creator in task.get("creators", []):
        row_index = int(creator["rowIndex"])
        result = by_row.get(row_index) or by_link.get(str(creator.get("pgyLink") or ""))
        if not result:
            continue
        for column in task.get("columns", []):
            column_index = int(column["index"]) + 1
            mapping = column.get("mapping")
            if mapping and mapping.get("source") == "pgy":
                value = value_for_dynamic_column(column, result)
                if set_if_blank(sheet, row_index, column_index, numeric_or_blank(value)):
                    filled.append({
                        "rowIndex": row_index,
                        "column": column["letter"],
                        "field": mapping.get("name", ""),
                        "source": "pgy",
                    })
            elif mapping and mapping.get("source") == "formula":
                cell = sheet.cell(row=row_index, column=column_index)
                formula = formula_for(mapping.get("id", ""), row_index, by_field_id, task)
                if is_blank(cell.value) and formula:
                    cell.value = formula
                    filled.append({
                        "rowIndex": row_index,
                        "column": column["letter"],
                        "field": mapping.get("name", ""),
                        "source": "formula",
                    })
            elif not mapping and column.get("displayLabel"):
                decision = (result.get("llmFieldDecisions") or {}).get(column.get("key"), {})
                if decision.get("supported") and not is_blank(decision.get("value")):
                    if set_if_blank(
                        sheet,
                        row_index,
                        column_index,
                        numeric_or_blank(decision.get("value")),
                    ):
                        filled.append({
                            "rowIndex": row_index,
                            "column": column["letter"],
                            "field": column.get("pathLabel", ""),
                            "source": "deepseek",
                        })
                elif is_blank(sheet.cell(row=row_index, column=column_index).value):
                    left_blank.append({
                        "rowIndex": row_index,
                        "column": column["letter"],
                        "field": column.get("pathLabel", ""),
                        "reason": decision.get("reason") or "现有蒲公英数据不能证明",
                    })

    creator_rows = [int(item["rowIndex"]) for item in task.get("creators", [])]
    if not creator_rows:
        raise RuntimeError("动态任务中没有达人")
    for column in task.get("columns", []):
        for row_index in range(min(creator_rows), max(creator_rows) + 1):
            cell = sheet.cell(row=row_index, column=int(column["index"]) + 1)
            apply_dynamic_format(cell, column.get("mapping"))
            if column.get("displayLabel") in ("账号数据表现", "推荐理由", "内容方向"):
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                cell.alignment = alignment

    output = Path(args.output).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    formulas, errors = formulas_and_errors(sheet)
    render_preview_svg(sheet, args.preview)
    validation = {
        "outputPath": str(output),
        "previewPath": str(Path(args.preview).expanduser().resolve()),
        "sheetName": sheet.title,
        "creatorCount": len(creator_rows),
        "task": task,
        "filled": filled,
        "leftBlank": left_blank,
        "tableInspection": table_snapshot(sheet),
        "formulaInspection": formulas,
        "errorInspection": errors,
        "appendedColumns": appended_columns,
    }
    write_validation(args.validation, validation)
    print(json.dumps({
        "ok": True,
        "outputPath": str(output),
        "previewPath": validation["previewPath"],
        "validationPath": str(Path(args.validation).expanduser().resolve()),
        "creatorCount": len(creator_rows),
        "filledCount": len(filled),
        "blankCount": len(left_blank),
    }, ensure_ascii=False))


def dynamic_export(args):
    task = json.loads(Path(args.task).read_text(encoding="utf-8"))
    results = json.loads(Path(args.results).read_text(encoding="utf-8"))
    mapped_columns = [
        column for column in task.get("columns", [])
        if (column.get("mapping") or {}).get("source") in ("pgy", "formula")
    ]
    if not mapped_columns:
        raise RuntimeError("动态任务没有可导出的字段")

    fixed_columns = [
        {"fieldId": "__creator_name", "name": "博主名称", "group": "笔记主页"},
        {"fieldId": "__pgy_link", "name": "蒲公英链接", "group": "笔记主页"},
    ]
    seen = {item["fieldId"] for item in fixed_columns}
    output_columns = list(fixed_columns)
    for column in mapped_columns:
        mapping = column.get("mapping") or {}
        field_id = mapping.get("id", "")
        if not field_id or field_id in seen:
            continue
        if field_id in ("pgy.creator.name", "pgy.creator.pgy_profile"):
            continue
        seen.add(field_id)
        output_columns.append({
            "fieldId": field_id,
            "name": mapping.get("label") or column.get("displayLabel") or mapping.get("name") or field_id,
            "group": column.get("topLabel") or column.get("parentLabel") or "蒲公英数据",
            "column": column,
        })

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "达人表现数据"
    group_fill = PatternFill("solid", fgColor="5B4FD4")
    header_fill = PatternFill("solid", fgColor="F2F0FF")
    for index, column in enumerate(output_columns, start=1):
        group_cell = sheet.cell(row=1, column=index, value=excel_safe_text(column.get("group") or "蒲公英数据"))
        group_cell.font = Font(bold=True, color="FFFFFF")
        group_cell.fill = group_fill
        group_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        metric_cell = sheet.cell(row=2, column=index, value=excel_safe_text(column["name"]))
        metric_cell.font = Font(bold=True, color="201A4D")
        metric_cell.fill = header_fill
        metric_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(len(column["name"]) * 2 + 4, 14),
            34,
        )
    merge_start = 1
    for index in range(2, len(output_columns) + 2):
        previous_group = output_columns[index - 2].get("group") if index - 2 < len(output_columns) else None
        current_group = output_columns[index - 1].get("group") if index - 1 < len(output_columns) else None
        if index <= len(output_columns) and current_group == previous_group:
            continue
        if index - merge_start > 1:
            sheet.merge_cells(start_row=1, start_column=merge_start, end_row=1, end_column=index - 1)
        merge_start = index
    sheet.row_dimensions[1].height = 38
    sheet.row_dimensions[2].height = 34
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = (
        f"A2:{get_column_letter(len(output_columns))}{max(2, len(task.get('creators') or results) + 2)}"
    )

    by_row = {int(item["rowIndex"]): item for item in results if item.get("rowIndex")}
    by_link = {str(item.get("pgyLink") or ""): item for item in results}
    output_letter_by_field = {
        column["fieldId"]: get_column_letter(index)
        for index, column in enumerate(output_columns, start=1)
    }
    filled = []
    left_blank = []
    creators = task.get("creators") or results
    for output_row, creator in enumerate(creators, start=3):
        source_row = int(creator.get("rowIndex") or output_row)
        result = by_row.get(source_row) or by_link.get(str(creator.get("pgyLink") or "")) or creator
        for output_column, definition in enumerate(output_columns, start=1):
            field_id = definition["fieldId"]
            value = ""
            source = "system"
            if field_id == "__creator_name":
                value = (
                    (result.get("fields") or {}).get("达人名称")
                    or result.get("nickname")
                    or creator.get("nickname")
                    or ""
                )
            elif field_id == "__pgy_link":
                value = result.get("pgyLink") or creator.get("pgyLink") or ""
            else:
                column = definition["column"]
                mapping = column.get("mapping") or {}
                source = mapping.get("source") or "pgy"
                if source == "formula":
                    value = formula_for(field_id, output_row, output_letter_by_field, task)
                else:
                    value = value_for_dynamic_column(column, result)

            cell = sheet.cell(row=output_row, column=output_column)
            if isinstance(value, str) and value.startswith("=") and source == "formula":
                cell.value = value
            else:
                cell.value = numeric_or_blank(value)
            if field_id == "__pgy_link" and isinstance(value, str) and value.startswith("http"):
                cell.hyperlink = value
                cell.style = "Hyperlink"
            if field_id not in ("__creator_name", "__pgy_link"):
                mapping = (definition.get("column") or {}).get("mapping") or {}
                apply_dynamic_format(cell, mapping)
                if is_blank(value):
                    left_blank.append({
                        "rowIndex": source_row,
                        "column": get_column_letter(output_column),
                        "field": definition["name"],
                        "reason": "蒲公英未返回该字段或字段依赖不足",
                    })
                else:
                    filled.append({
                        "rowIndex": source_row,
                        "column": get_column_letter(output_column),
                        "field": definition["name"],
                        "source": source,
                    })

    audit_sheet = workbook.create_sheet("采集说明")
    audit_headers = ["博主名称", "蒲公英链接", "采集状态", "未完成原因"]
    for column_index, header in enumerate(audit_headers, start=1):
        cell = audit_sheet.cell(row=1, column=column_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = group_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index, creator in enumerate(creators, start=2):
        source_row = int(creator.get("rowIndex") or row_index)
        result = by_row.get(source_row) or by_link.get(str(creator.get("pgyLink") or "")) or creator
        audit_sheet.cell(
            row=row_index,
            column=1,
            value=excel_safe_text(
                (result.get("fields") or {}).get("达人名称")
                or result.get("nickname")
                or creator.get("nickname")
                or ""
            ),
        )
        link = result.get("pgyLink") or creator.get("pgyLink") or ""
        link_cell = audit_sheet.cell(row=row_index, column=2, value=excel_safe_text(link))
        if isinstance(link, str) and link.startswith("http"):
            link_cell.hyperlink = link
            link_cell.style = "Hyperlink"
        audit_sheet.cell(row=row_index, column=3, value=excel_safe_text(result.get("dataStatus") or ""))
        audit_sheet.cell(row=row_index, column=4, value=excel_safe_text(result.get("error") or ""))
    audit_sheet.freeze_panes = "A2"
    audit_sheet.auto_filter.ref = f"A1:D{max(1, len(creators) + 1)}"
    for column_letter, width in {"A": 20, "B": 48, "C": 16, "D": 54}.items():
        audit_sheet.column_dimensions[column_letter].width = width

    output = Path(args.output).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    formulas, errors = formulas_and_errors(sheet)
    render_preview_svg(sheet, args.preview)
    validation = {
        "outputPath": str(output),
        "previewPath": str(Path(args.preview).expanduser().resolve()),
        "sheetName": sheet.title,
        "creatorCount": len(creators),
        "task": task,
        "filled": filled,
        "leftBlank": left_blank,
        "tableInspection": table_snapshot(sheet),
        "formulaInspection": formulas,
        "errorInspection": errors,
    }
    write_validation(args.validation, validation)
    print(json.dumps({
        "ok": True,
        "outputPath": str(output),
        "previewPath": validation["previewPath"],
        "validationPath": str(Path(args.validation).expanduser().resolve()),
        "creatorCount": len(creators),
        "filledCount": len(filled),
        "blankCount": len(left_blank),
    }, ensure_ascii=False))


def normalized_sheet_name(value):
    return re.sub(r"\s+", "", str(value or "")).lower()


def medela_sheet(workbook):
    for sheet in workbook.worksheets:
        if "kol提报表" in normalized_sheet_name(sheet.title):
            return sheet
    return workbook.worksheets[0]


def creator_id_from_url(url):
    match = re.search(r"/blogger-detail/([^/?#]+)", str(url or ""))
    return match.group(1) if match else ""


def find_header_column(sheet, label, max_rows=6):
    for row_index in range(1, min(sheet.max_row, max_rows) + 1):
        for column_index in range(1, sheet.max_column + 1):
            value = re.sub(r"\s+", "", str(sheet.cell(row=row_index, column=column_index).value or ""))
            if label in value:
                return row_index, column_index
    return None, None


def medela_creators(sheet):
    header_row, link_column = find_header_column(sheet, "蒲公英链接")
    if not link_column:
        link_column = 9
        header_row = 2
    _, sequence_column = find_header_column(sheet, "序号")
    _, creator_type_column = find_header_column(sheet, "达人类型")
    _, nickname_column = find_header_column(sheet, "名称")
    sequence_column = sequence_column or 2
    creator_type_column = creator_type_column or 3
    nickname_column = nickname_column or 4
    creators = []
    for row_index in range((header_row or 2) + 1, sheet.max_row + 1):
        pgy_link = str(cell_value(sheet.cell(row=row_index, column=link_column)) or "").strip()
        if "pgy.xiaohongshu.com" not in pgy_link:
            continue
        creators.append({
            "rowIndex": row_index,
            "sequence": json_value(sheet.cell(row=row_index, column=sequence_column).value) or "",
            "creatorType": str(sheet.cell(row=row_index, column=creator_type_column).value or "").strip(),
            "nickname": str(sheet.cell(row=row_index, column=nickname_column).value or "").strip(),
            "pgyLink": pgy_link,
            "creatorId": creator_id_from_url(pgy_link),
        })
    return creators


def medela_extract(args):
    workbook = load_book(args.input)
    sheet = medela_sheet(workbook)
    creators = medela_creators(sheet)
    if not creators:
        raise RuntimeError("KOL提报表中未找到蒲公英达人链接")
    output = Path(args.creators).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(creators, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "ok": True,
        "count": len(creators),
        "sheetName": sheet.title,
    }, ensure_ascii=False))


def medela_fill(args):
    results = json.loads(Path(args.results).read_text(encoding="utf-8"))
    workbook = load_book(args.input)
    sheet = medela_sheet(workbook)
    creators = medela_creators(sheet)
    if not creators:
        raise RuntimeError("KOL提报表中未找到蒲公英达人链接")
    by_row = {int(item["rowIndex"]): item for item in results if item.get("rowIndex")}
    by_link = {str(item.get("pgyLink") or ""): item for item in results}
    updated_cells = []
    field_columns = [
        ("小红书主页链接", 10),
        ("粉丝数（w）", 11),
        ("合作笔记曝光中位数（90天）", 12),
        ("曝光来源-发现页", 13),
        ("曝光来源-搜索页", 14),
        ("曝光来源-关注页", 15),
        ("曝光来源-博主个人页", 16),
        ("曝光来源-附近页", 17),
        ("曝光来源-其他", 18),
        ("合作笔记阅读中位数（90天）", 19),
        ("预估合作笔记自然流曝光（90天）", 20),
        ("预估合作笔记自然流阅读（90天）", 21),
        ("报价", 22),
    ]
    for creator in creators:
        row = int(creator["rowIndex"])
        result = by_row.get(row) or by_link.get(creator["pgyLink"])
        if not result:
            continue
        fields = result.get("fields") or {}
        sheet.cell(row=row, column=5).value = excel_safe_text(result.get("recommendationReason") or "")
        updated_cells.append(f"E{row}")
        for field, column in field_columns:
            value = fields.get(field, "")
            sheet.cell(row=row, column=column).value = (
                excel_safe_text(value or "") if field == "小红书主页链接" else numeric_or_blank(value)
            )
            updated_cells.append(f"{get_column_letter(column)}{row}")
        sheet.cell(row=row, column=23).value = f'=IF(V{row}="","",ROUND(V{row}*1.1,2))'
        sheet.cell(row=row, column=24).value = f'=IF(W{row}="","",ROUND(W{row}*1.02,2))'
        sheet.cell(row=row, column=25).value = f'=IF(U{row}>0,V{row}/U{row},"")'
        updated_cells.extend([f"W{row}", f"X{row}", f"Y{row}"])

    rows = [int(item["rowIndex"]) for item in creators]
    for row in range(min(rows), max(rows) + 1):
        recommendation_alignment = copy(sheet.cell(row=row, column=5).alignment)
        recommendation_alignment.wrap_text = True
        sheet.cell(row=row, column=5).alignment = recommendation_alignment
        link_alignment = copy(sheet.cell(row=row, column=10).alignment)
        link_alignment.wrap_text = True
        sheet.cell(row=row, column=10).alignment = link_alignment
        sheet.cell(row=row, column=11).number_format = "0.0000"
        sheet.cell(row=row, column=12).number_format = "#,##0"
        for column in range(13, 19):
            sheet.cell(row=row, column=column).number_format = "0.0%"
        for column in range(19, 22):
            sheet.cell(row=row, column=column).number_format = "#,##0"
        for column in range(22, 25):
            sheet.cell(row=row, column=column).number_format = "¥#,##0.00"
        sheet.cell(row=row, column=25).number_format = "0.00"

    output = Path(args.output).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    formulas, errors = formulas_and_errors(sheet)
    render_preview_svg(sheet, args.preview, max_rows=max(rows), max_cols=26)
    validation = {
        "outputPath": str(output),
        "previewPath": str(Path(args.preview).expanduser().resolve()),
        "sheetName": sheet.title,
        "creatorCount": len(creators),
        "updatedRows": rows,
        "updatedCells": updated_cells,
        "keyInspection": table_snapshot(sheet, max_rows=max(rows), max_cols=26),
        "formulaInspection": formulas,
        "errorInspection": errors,
    }
    write_validation(args.validation, validation)
    print(json.dumps({
        "ok": True,
        "outputPath": str(output),
        "previewPath": validation["previewPath"],
        "validationPath": str(Path(args.validation).expanduser().resolve()),
        "creatorCount": len(creators),
    }, ensure_ascii=False))


def parser():
    root = argparse.ArgumentParser(description="寻达生产环境 Excel 读写桥")
    commands = root.add_subparsers(dest="command", required=True)

    inspect_parser = commands.add_parser("inspect")
    inspect_parser.add_argument("--input", required=True)
    inspect_parser.set_defaults(func=inspect_workbook)

    attach_links_parser = commands.add_parser("attach-links")
    attach_links_parser.add_argument("--input", required=True)
    attach_links_parser.add_argument("--output")
    attach_links_parser.add_argument("--links", required=True)
    attach_links_parser.set_defaults(func=attach_links)

    dynamic_parser = commands.add_parser("dynamic-fill")
    dynamic_parser.add_argument("--input", required=True)
    dynamic_parser.add_argument("--task", required=True)
    dynamic_parser.add_argument("--results", required=True)
    dynamic_parser.add_argument("--output", required=True)
    dynamic_parser.add_argument("--preview", required=True)
    dynamic_parser.add_argument("--validation", required=True)
    dynamic_parser.set_defaults(func=dynamic_fill)

    dynamic_export_parser = commands.add_parser("dynamic-export")
    dynamic_export_parser.add_argument("--task", required=True)
    dynamic_export_parser.add_argument("--results", required=True)
    dynamic_export_parser.add_argument("--output", required=True)
    dynamic_export_parser.add_argument("--preview", required=True)
    dynamic_export_parser.add_argument("--validation", required=True)
    dynamic_export_parser.set_defaults(func=dynamic_export)

    medela_extract_parser = commands.add_parser("medela-extract")
    medela_extract_parser.add_argument("--input", required=True)
    medela_extract_parser.add_argument("--creators", required=True)
    medela_extract_parser.set_defaults(func=medela_extract)

    medela_fill_parser = commands.add_parser("medela-fill")
    medela_fill_parser.add_argument("--input", required=True)
    medela_fill_parser.add_argument("--results", required=True)
    medela_fill_parser.add_argument("--output", required=True)
    medela_fill_parser.add_argument("--preview", required=True)
    medela_fill_parser.add_argument("--validation", required=True)
    medela_fill_parser.set_defaults(func=medela_fill)
    return root


if __name__ == "__main__":
    arguments = parser().parse_args()
    arguments.func(arguments)
